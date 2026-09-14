"""Full regression against the client-approved July 2026 output workbook."""
from pathlib import Path
import sys
import datetime as dt
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.engine import readers
from app.engine.golden import calculate

REF = ROOT / "reference"


def expected_rows():
    wb = openpyxl.load_workbook(REF / "80-20_Table_20260903_120741.xlsx",
                                read_only=True, data_only=False)
    ws = wb["80-20 Table"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
    review_ws = wb["Needs Review"]
    review_headers = [c.value for c in next(review_ws.iter_rows(min_row=1, max_row=1))]
    review = [dict(zip(review_headers, values))
              for values in review_ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return rows, review


def actual_rows():
    result = calculate(
        readers.read_daybook(REF / "DayBookRegister.xlsx"),
        readers.read_voucher(REF / "Search Voucher.xlsx"),
        readers.read_creditors(REF / "Creditors Details (1).xlsx"),
        {"include_fixed_groups": ["Current Assets", "Fixed Assets"],
         "exclude_parent_groups": ["TDS Receivable"]},
    )
    return result["rows"], result["exceptions"]


FIELD_MAP = {
    "Unique_Voucher_ID": "voucher_id", "Head": "head", "Year": "year",
    "Voucher_Type": "voucher_type", "Voucher_No": "voucher_no", "Bill_No": "bill_no",
    "Bill_Date": "bill_date", "Project Name": "project", "Narration": "narration",
    "Account_Head": "account_head", "Account_Name": "account_name",
    "Account Ledger": "ledger", "GST No.": "gstin", "Debit": "debit",
    "Credit": "credit", "Closing": "closing", "Formula Key": "formula_key",
    "80-20": "eligibility", "GST Status": "gst_status",
}


def normal(value):
    if value is None:
        return ""
    return value.date() if isinstance(value, dt.datetime) else value


def main():
    expected, expected_review = expected_rows()
    actual, actual_review = actual_rows()
    assert len(actual) == len(expected) == 1232, (len(actual), len(expected))
    for index, (want, got) in enumerate(zip(expected, actual), start=2):
        for excel_name, app_name in FIELD_MAP.items():
            left, right = normal(want[excel_name]), normal(got[app_name])
            if left != right:
                raise AssertionError(
                    f"row {index} {excel_name}: expected {left!r}, got {right!r}")
        expected_month = want["Month"].strftime("%b-%y")
        assert got["month"] == expected_month, (index, expected_month, got["month"])
        assert got["voucher_date"] == want["Voucher_Date"].date()

    assert len(actual_review) == len(expected_review) == 20
    # Diagnostics and source narration are an app-only review layer.  They
    # must never alter the approved table or the script's Needs Review list.
    assert sum(bool(r["source_narration"]) for r in actual) == 1232
    assert sum("gstin_bad_checksum" in r["flags"] for r in actual) == 3
    assert not any("spec_" in r["flags"] for r in actual)
    for want, got in zip(expected_review, actual_review):
        assert got["voucher_id"] == want["Unique_Voucher_ID"]
        assert got["ledger"] == (want["Account Ledger (resolved)"] or "")
        assert got["reason"] == (want["Issue"] or "")
        assert got["suggestions"] == (want["Suggestions"] or "")
    print("PASS: all 1,232 detail rows and 20 review items match the golden workbook")


if __name__ == "__main__":
    main()
