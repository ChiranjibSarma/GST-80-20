"""End-to-end HTTP workflow: upload golden inputs, persist, freeze, reject change."""
import os
from pathlib import Path
import tempfile
import io
import re
import html
import datetime as dt
import json
import sqlite3
import openpyxl

TEST_VAR = tempfile.mkdtemp(prefix="gst8020-test-")
os.environ["VAR_DIR"] = TEST_VAR
os.environ["BOOTSTRAP_ADMIN_EMAIL"] = "admin@test.local"
os.environ["BOOTSTRAP_ADMIN_PASSWORD"] = "GoldenTestPassword1"

from fastapi.testclient import TestClient
from sqlalchemy import select, func

from app.main import app
from app import license as licensing
from app.db import SessionLocal
from app.config import BACKUP_DIR
from app.models import Run, RunRow, Rectification, AuditLog, Creditor, LicenseActivation

# Pre-signed test fixtures use a different public key from the client package.
# No private key or signing implementation is distributed with the tests.
FIXTURES = Path(__file__).resolve().parent / "fixtures"
licensing.PUBLIC_KEY_PATH = Path(TEST_VAR) / "public.pem"
licensing.PUBLIC_KEY_PATH.write_bytes((FIXTURES / "license_public_key.pem").read_bytes())
test_document = json.loads((FIXTURES / "license_valid.json").read_text(encoding="utf-8"))
test_payload = test_document["payload"]
test_installation_id = test_payload["installation_id"]
licensing.INSTALLATION_ID_FILE.write_text(test_installation_id, encoding="ascii")
licensing.LICENSE_FILE.write_text(json.dumps(test_document), encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "reference"


def upload_payload():
    return {
        "daybook": ("DayBookRegister.xlsx", (REF / "DayBookRegister.xlsx").read_bytes()),
        "voucher": ("Search Voucher.xlsx", (REF / "Search Voucher.xlsx").read_bytes()),
        "creditors": ("Creditors Details (1).xlsx", (REF / "Creditors Details (1).xlsx").read_bytes()),
    }


def main():
    with TestClient(app) as client:
        login_page = client.get("/login")
        assert login_page.status_code == 200
        csrf = html.unescape(re.search(r'name="csrf_token_value" value="([^"]+)"',
                                        login_page.text).group(1))
        assert client.post("/login", data={"email": "admin@test.local",
                                           "password": "GoldenTestPassword1"}).status_code == 403
        response = client.post("/login", data={"email": "admin@test.local",
                                                "password": "GoldenTestPassword1",
                                                "next": "/", "csrf_token_value": csrf},
                               follow_redirects=False)
        assert response.status_code == 303
        form = client.get("/gst8020/new")
        assert form.status_code == 200
        for filename in ("Day_Book_Register_Template.xlsx", "Search_Voucher_Template.xlsx",
                         "Creditors_Details_Template.xlsx"):
            assert filename in form.text
            template = client.get(f"/static/input-templates/{filename}")
            assert template.status_code == 200 and template.content.startswith(b"PK")
        response = client.post("/gst8020/new", data={"label": "July 2026 golden run",
                                                    "csrf_token_value": csrf},
                               files=upload_payload(), follow_redirects=False)
        assert response.status_code == 303, response.text[:1000]
        run_id = int(response.headers["location"].rstrip("/").split("/")[-1])
        backups = list(BACKUP_DIR.glob("gst8020-*.sqlite3"))
        assert len(backups) == 1
        with sqlite3.connect(backups[0]) as backup_db:
            assert backup_db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert backup_db.execute("SELECT count(*) FROM run_rows").fetchone()[0] == 1232
        with SessionLocal() as db:
            assert db.scalar(select(func.count(RunRow.id)).where(RunRow.run_id == run_id)) == 1232
            assert db.scalar(select(func.count(Rectification.id)).where(Rectification.run_id == run_id)) == 20
            assert db.scalar(select(func.count(RunRow.id)).where(
                RunRow.run_id == run_id, RunRow.narration != "")) == 1232
            assert db.scalar(select(func.count(RunRow.id)).where(
                RunRow.run_id == run_id, RunRow.flags.like("%gstin_bad_checksum%"))) == 3
            review_id = db.scalar(select(Rectification.id).where(Rectification.run_id == run_id))
            creditor_count = db.scalar(select(func.count(Creditor.id)))

        exceptions = client.get(f"/gst8020/runs/{run_id}/exceptions")
        assert exceptions.status_code == 200
        assert "Email rule differs" not in exceptions.text
        assert "GSTIN failed validation" in exceptions.text
        detail_page = client.get(f"/gst8020/runs/{run_id}/rows")
        assert detail_page.status_code == 200
        assert "Source narration" in detail_page.text
        assert "Correct for next run" not in detail_page.text
        assert client.post("/gst8020/overrides", data={
            "raw": "Vendor A", "ledger": "Manual", "csrf_token_value": csrf,
        }).status_code == 409
        assert client.post("/gst8020/masters", data={
            "ineligible_keywords": "", "csrf_token_value": csrf,
        }).status_code == 409

        exported = client.get(f"/gst8020/runs/{run_id}/export.xlsx")
        assert exported.status_code == 200
        got_wb = openpyxl.load_workbook(io.BytesIO(exported.content), read_only=True, data_only=False)
        want_wb = openpyxl.load_workbook(REF / "80-20_Table_20260903_120741.xlsx",
                                         read_only=True, data_only=False)
        got_rows = list(got_wb["80-20 Table"].iter_rows(values_only=True))
        want_rows = list(want_wb["80-20 Table"].iter_rows(values_only=True))
        assert got_rows == want_rows
        got_review = list(got_wb["Needs Review"].iter_rows(values_only=True))
        want_review = list(want_wb["Needs Review"].iter_rows(values_only=True))
        assert got_review == want_review
        for tab in ("Report", "Instant Review Pivot"):
            got = list(got_wb[tab].iter_rows(values_only=True))
            want = list(want_wb[tab].iter_rows(values_only=True))
            assert got == want, f"{tab} differs"
        quality = list(got_wb["Data Quality"].iter_rows(values_only=True))
        assert quality[0][-2:] == ("Source Narration", "Flags")
        assert any("gstin_bad_checksum" in str(row[-1]) and row[-2] for row in quality[1:])
        got_wb.close()
        want_wb.close()

        response = client.post(f"/gst8020/runs/{run_id}/freeze",
                               data={"csrf_token_value": csrf}, follow_redirects=False)
        assert response.status_code == 303
        with SessionLocal() as db:
            run = db.get(Run, run_id)
            assert run.status == "frozen" and run.frozen_at is not None
            assert db.scalar(select(func.count(AuditLog.id)).where(
                AuditLog.action == "freeze_run", AuditLog.entity_id == str(run_id))) == 1

        response = client.post("/gst8020/new", data={"label": "Forbidden replacement",
                                                    "csrf_token_value": csrf},
                               files=upload_payload(), follow_redirects=False)
        assert response.status_code == 409
        assert "is frozen" in response.text
        with SessionLocal() as db:
            assert db.scalar(select(func.count(Run.id))) == 1
            assert db.scalar(select(func.count(Creditor.id))) == creditor_count
        response = client.post(f"/gst8020/rectifications/{review_id}",
                               data={"status": "wont_fix", "csrf_token_value": csrf},
                               follow_redirects=False)
        assert response.status_code == 409

        with SessionLocal() as db:
            activation = db.get(LicenseActivation, test_payload["license_id"])
            expiry = licensing.as_utc(activation.activated_at) + dt.timedelta(days=14)
        licensing.utcnow = lambda: expiry
        assert client.post("/admin/users", data={
            "name": "Blocked", "email": "blocked@test.local", "password": "Password12345",
            "role": "viewer", "csrf_token_value": csrf,
        }).status_code == 423
        assert client.get(f"/gst8020/runs/{run_id}/export.xlsx").status_code == 200

    print("PASS: upload/export match golden workbook; freeze and licence expiry reject writes")


if __name__ == "__main__":
    main()
