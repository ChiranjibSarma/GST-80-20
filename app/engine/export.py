"""Excel exports: the golden-compatible workbook and GSTIN completion sheet."""
import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from ..models import RunRow, Rectification, Creditor
from .calc import summarise, summarise_by_project, vendor_concentration, merged_config

HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")


def _sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR_FONT, HDR_FILL, Alignment(vertical="center")
    for r in rows:
        ws.append([r.get(c, "") for c in columns])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, len(str(c)) + 3), 44)
    return ws


COLUMNS = ["Unique_Voucher_ID", "Head", "Year", "Month", "Voucher_Date", "Voucher_Type",
           "Voucher_No", "Bill_No", "Bill_Date", "Project Name", "Narration", "Account_Head",
           "Account_Name", "Account Ledger", "GST No.", "Debit", "Credit", "Closing",
           "Formula Key", "80-20", "GST Status"]


def build_workbook(db, run) -> io.BytesIO:
    recorded_cfg = json.loads(run.config_snapshot or "{}")
    cfg = merged_config(recorded_cfg)
    # The approved script preserves Day Book line order; RunRow IDs capture it.
    db_rows = list(db.scalars(select(RunRow).where(RunRow.run_id == run.id)
                              .order_by(RunRow.id)))
    rows = [{
        "Unique_Voucher_ID": r.voucher_id, "Head": r.head, "Year": r.year,
        "Month": r.voucher_date.replace(day=1) if r.voucher_date else None,
        "Voucher_Date": r.voucher_date, "Voucher_Type": r.voucher_type,
        "Voucher_No": int(r.voucher_no) if str(r.voucher_no).isdigit() else r.voucher_no,
        "Bill_No": r.bill_no, "Bill_Date": r.bill_date, "Project Name": r.project,
        # The approved Python leaves this column blank.  Source narration is
        # retained in the app for investigation, not silently put into parity export.
        "Narration": "", "Account_Head": r.account_head, "Account_Name": r.account_name,
        "Account Ledger": r.ledger, "GST No.": r.gstin, "Debit": r.debit, "Credit": r.credit,
        "Closing": r.closing, "Formula Key": r.formula_key, "80-20": r.eligibility,
        "GST Status": r.gst_status,
    } for r in db_rows]
    plain = [dict(project=r.project, eligibility=r.eligibility, gst_status=r.gst_status,
                  closing=r.closing, head=r.head, ledger=r.ledger, gstin=r.gstin,
                  voucher_id=r.voucher_id) for r in db_rows]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    detail = _sheet(wb, "80-20 Table", COLUMNS, rows)
    for cell in detail["D"][1:]:
        cell.number_format = "mmm-yy"
    for column in ("E", "I"):
        for cell in detail[column][1:]:
            cell.number_format = "dd-mmm-yy"
    for column in ("P", "Q", "R"):
        for cell in detail[column][1:]:
            cell.number_format = "#,##0.00"

    # -- approved Report: the same live SUMIFS and absolute-value GST mix --
    ws = wb.create_sheet("Report")
    for col, heading in enumerate(["Head", "80-20", "GST Status", "Sum of Closing"], 1):
        ws.cell(1, col, heading)
    last = len(rows) + 1
    closing_rng = f"'80-20 Table'!$R$2:$R${last}"
    head_rng = f"'80-20 Table'!$B$2:$B${last}"
    elig_rng = f"'80-20 Table'!$T$2:$T${last}"
    gst_rng = f"'80-20 Table'!$U$2:$U${last}"
    heads = sorted({r["Head"] for r in rows}) or ["WIP"]
    row_no, grand_total_cells = 2, []
    for head in heads:
        head_totals = []
        for eligibility in ("Eligible", "Ineligible"):
            block_start = row_no
            for gst_status in ("Registered", "Unregistered"):
                ws.cell(row_no, 1, head if eligibility == "Eligible" and gst_status == "Registered" else None)
                ws.cell(row_no, 2, eligibility if gst_status == "Registered" else None)
                ws.cell(row_no, 3, gst_status)
                ws.cell(row_no, 4,
                    f'=SUMIFS({closing_rng},{head_rng},"{head}",{elig_rng},"{eligibility}",{gst_rng},"{gst_status}")')
                row_no += 1
            ws.cell(row_no, 2, f"{eligibility} Total")
            ws.cell(row_no, 4, f"=SUM(D{block_start}:D{row_no - 1})")
            head_totals.append(row_no)
            row_no += 1
        ws.cell(row_no, 1, f"{head} Total")
        ws.cell(row_no, 4, f"=D{head_totals[0]}+D{head_totals[1]}")
        grand_total_cells.append(row_no)
        row_no += 1
    ws.cell(row_no, 1, "Grand Total")
    ws.cell(row_no, 4, "=" + "+".join(f"D{x}" for x in grand_total_cells))
    for col, heading in enumerate(["Eligible GST Mix", "Eligible Amount", "% of Eligible"], 6):
        ws.cell(1, col, heading)
    for row_no, label in ((2, "Registered"), (3, "Unregistered"), (4, "Eligible Total")):
        ws.cell(row_no, 6, label)
    ws.cell(2, 7, f'=SUMIFS({closing_rng},{elig_rng},"Eligible",{gst_rng},"Registered")')
    ws.cell(3, 7, f'=SUMIFS({closing_rng},{elig_rng},"Eligible",{gst_rng},"Unregistered")')
    ws.cell(4, 7, "=SUM(G2:G3)")
    for row_no in (2, 3):
        ws.cell(row_no, 8, f'=IFERROR(ABS(G{row_no})/SUM(ABS($G$2),ABS($G$3)),0)')
        ws.cell(row_no, 8).number_format = "0.00%"
    for col in (1, 2, 3, 4, 6, 7, 8):
        ws.cell(1, col).font, ws.cell(1, col).fill = HDR_FONT, HDR_FILL
    for col, width in enumerate([18, 14, 26, 16, 3, 24, 18, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # -- approved Instant Review Pivot --------------------------------------
    pivot = wb.create_sheet("Instant Review Pivot")
    pivot.append(["Head", "80-20", "GST Status", "Line Count", "Sum of Closing",
                  "% of Eligible Amount"])
    for cell in pivot[1]:
        cell.font, cell.fill = HDR_FONT, HDR_FILL
    for head in heads:
        for eligibility in ("Eligible", "Ineligible"):
            for status in ("Registered", "Unregistered"):
                matches = [r for r in rows if r["Head"] == head and r["80-20"] == eligibility
                           and r["GST Status"] == status]
                if not matches:
                    continue
                amount = sum(r["Closing"] or 0 for r in matches)
                denominator = sum(abs(r["Closing"] or 0) for r in rows
                                  if r["Head"] == head and r["80-20"] == "Eligible")
                percentage = (abs(amount) / denominator if denominator else 0) if eligibility == "Eligible" else None
                pivot.append([head, eligibility, status, len(matches), amount, percentage])
    pivot.freeze_panes = "A2"
    pivot.auto_filter.ref = f"A1:F{pivot.max_row}"
    for row_no in range(2, pivot.max_row + 1):
        pivot.cell(row_no, 5).number_format = "#,##0.00"
        pivot.cell(row_no, 6).number_format = "0.00%"
    for col, width in enumerate([18, 14, 52, 12, 18, 20], 1):
        pivot.column_dimensions[get_column_letter(col)].width = width

    # -- per project ----------------------------------------------------------
    _sheet(wb, "By Project",
           ["Project", "Eligible", "Registered", "Unregistered",
            "Registered %", "Shortfall", "Tax exposure"],
           [{"Project": p, "Eligible": round(v["eligible"], 2),
             "Registered": round(v["registered"], 2),
             "Unregistered": round(v["unregistered"], 2),
             "Registered %": round(v["pct"], 2), "Shortfall": round(v["gap"], 2),
             "Tax exposure": round(v["exposure"], 2)}
            for p, v in summarise_by_project(plain, cfg).items()])

    # -- rectification queue --------------------------------------------------
    items = list(db.scalars(select(Rectification).where(Rectification.run_id == run.id)
                            .order_by(Rectification.id)))
    review_sheet = _sheet(wb, "Needs Review",
           ["Unique_Voucher_ID", "Account Ledger (resolved)", "Issue", "Suggestions"],
           [{"Unique_Voucher_ID": i.voucher_id, "Account Ledger (resolved)": i.ledger,
             "Issue": i.reason, "Suggestions": i.suggestions}
            for i in items])
    wb.move_sheet(review_sheet, offset=-1)

    # -- vendor concentration --------------------------------------------------
    _sheet(wb, "Unregistered Vendors",
           ["Supplier (grouped)", "Amount", "Name variants", "Vouchers", "All spellings"],
           [{"Supplier (grouped)": c["cluster"], "Amount": round(c["amount"], 2),
             "Name variants": c["variants"], "Vouchers": c["vouchers"],
             "All spellings": " | ".join(c["names"])}
            for c in vendor_concentration(plain, limit=200)])

    # -- data quality -----------------------------------------------------------
    _sheet(wb, "Data Quality",
           ["Unique_Voucher_ID", "Account Ledger", "Account_Name", "Closing",
            "Source Narration", "Flags"],
           [{"Unique_Voucher_ID": r.voucher_id, "Account Ledger": r.ledger,
             "Account_Name": r.account_name, "Closing": round(r.closing, 2),
             "Source Narration": r.narration, "Flags": r.flags}
            for r in db_rows if r.flags])

    # -- masters as used --------------------------------------------------------
    ws = wb.create_sheet("Masters Used")
    ws.append(["Setting", "Value as used for this run"])
    for c in ws[1]:
        c.font, c.fill = HDR_FONT, HDR_FILL
    ws.append(["Run", f"{run.label} ({run.period_month}, {run.financial_year}) - {run.status}"])
    ws.append(["Source files", json.dumps(json.loads(run.source_files or "{}"))])
    # Show the configuration actually recorded with this run. Older frozen
    # snapshots may contain controls which are no longer accepted for new runs.
    for k, v in recorded_cfg.items():
        ws.append([k, ", ".join(map(str, v)) if isinstance(v, list)
                   else json.dumps(v) if isinstance(v, dict) else str(v)])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def missing_gstin_workbook(db) -> io.BytesIO:
    rows = list(db.scalars(select(Creditor)
                           .where((Creditor.gstin == "") | (Creditor.gstin_valid != "ok"))
                           .order_by(Creditor.account_name)))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Complete GSTIN"
    ws.append(["Account_Name", "Group_Name", "PAN_no", "GSTIN", "Current issue"])
    for c in ws[1]:
        c.font, c.fill = HDR_FONT, HDR_FILL
    for r in rows:
        ws.append([r.account_name, r.group_name, r.pan, r.gstin,
                   "No GSTIN on file" if not r.gstin else f"Invalid: {r.gstin_valid}"])
    ws.freeze_panes = "A2"
    for i, w in enumerate([46, 26, 16, 20, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    note = wb.create_sheet("How to use this")
    for line in [
        "Fill in the GSTIN column for any supplier you can confirm, then upload this file",
        "on the Creditors page of the GST 80:20 module.",
        "",
        "Leave a row blank if the supplier is genuinely unregistered - that is a valid answer",
        "and the calculation treats it correctly.",
        "",
        "Every GSTIN is checked for structure and checksum on upload. Anything that fails is",
        "reported back to you and not saved, so a mistyped number cannot enter the master.",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 96

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
