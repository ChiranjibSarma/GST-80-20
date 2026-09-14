"""Readers for the three Tally exports.

Each export carries a banner block above the real header row, and the header
row is located by the columns it must contain rather than by a fixed offset,
so a change in the banner does not break ingestion.
"""
import datetime as dt
import re
import openpyxl

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def key(v) -> str:
    return norm(v).upper()


def num(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip() or 0)
    except ValueError:
        return 0.0


def parse_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = norm(v)
    if not s:
        return None
    for f in ("%d %b %Y", "%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


class IngestError(Exception):
    """Raised with a message written for the person who uploaded the file."""


def read_sheet(path_or_file, must_have, label):
    try:
        wb = openpyxl.load_workbook(path_or_file, read_only=True, data_only=True)
    except Exception as exc:
        # A corrupt, encrypted or legacy .xls file must become a user-facing
        # upload error, never an unhandled server exception.
        raise IngestError(f"{label}: could not read an .xlsx workbook ({type(exc).__name__}).") from exc
    ws = wb.worksheets[0]
    header_row, header = None, None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=60, values_only=True), 1):
        present = {key(c) for c in row if c is not None}
        if all(m in present for m in must_have):
            header_row, header = i, [norm(c) for c in row]
            break
    if header_row is None:
        wb.close()
        raise IngestError(
            f"{label}: could not find the column headings. This file needs the columns "
            f"{', '.join(must_have)} — check that you exported the right report from Tally.")
    # Tally may vary heading case/spacing.  Normalize to the exact field names
    # consumed by the approved Python script, rather than accepting a heading
    # and then silently returning an empty field under another spelling.
    canonical = {
        "DATE": "Date", "BILL_DATE": "Bill_Date", "VOUCHER_DATE": "Voucher_Date",
        "VOUCHER_TYPE": "Voucher_Type", "VOUCHER_NO": "Voucher_No",
        "PARTICULARS": "Particulars", "FIXED_GROUP_NAME": "Fixed_Group_Name",
        "PARENT_GROUP_NAME": "Parent_Group_Name", "DEBIT": "Debit", "CREDIT": "Credit",
        "BILL_NO": "Bill_No", "COST_CENTRE": "Cost_Centre", "NARRATION": "Narration",
        "DEBIT_AMOUNT": "Debit_Amount", "CREDIT_AMOUNT": "Credit_Amount",
        "STATUS": "Status", "ACCOUNT_NAME": "Account_Name", "GROUP_NAME": "Group_Name",
        "GSTIN": "GSTIN", "PAN_NO": "PAN_no", "LEGAL_CHEQUE_NAME": "Legal_Cheque_Name",
    }
    idx = {}
    for i, heading in enumerate(header):
        if not heading:
            continue
        field = canonical.get(key(heading), heading)
        if field in idx:
            wb.close()
            raise IngestError(f"{label}: duplicate column heading {field!r}.")
        idx[field] = i
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(c not in (None, "") for c in row):
            continue
        rows.append({h: (row[i] if i < len(row) else None) for h, i in idx.items()})
    wb.close()
    if not rows:
        raise IngestError(f"{label}: the file has headings but no data rows.")
    return rows


DAYBOOK_COLS = ["DATE", "VOUCHER_TYPE", "VOUCHER_NO", "PARTICULARS", "FIXED_GROUP_NAME",
                "PARENT_GROUP_NAME", "DEBIT", "CREDIT", "BILL_NO", "BILL_DATE",
                "COST_CENTRE", "NARRATION"]
VOUCHER_COLS = ["VOUCHER_DATE", "VOUCHER_NO", "VOUCHER_TYPE", "PARTICULARS",
                "DEBIT_AMOUNT", "CREDIT_AMOUNT", "STATUS"]
CREDITOR_COLS = ["ACCOUNT_NAME", "GROUP_NAME", "GSTIN"]


def read_daybook(f):   return read_sheet(f, DAYBOOK_COLS, "Day Book Register")
def read_voucher(f):   return read_sheet(f, VOUCHER_COLS, "Search Voucher")
def read_creditors(f): return read_sheet(f, CREDITOR_COLS, "Creditors Details")
